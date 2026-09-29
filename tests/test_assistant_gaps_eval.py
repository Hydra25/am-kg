from src.analysis.gaps import find_gaps
from src.evaluation.metrics import evaluate, graph_quality
from src.knowledge_graph.kb import KnowledgeBase
from src.llm.client import FakeLLM
from src.normalization.entity_normalizer import EntityNormalizer
from src.retrieval.assistant import ResearchAssistant


def test_assistant_graph_only(demo_corpus):
    kb = KnowledgeBase.load(demo_corpus)
    a = ResearchAssistant(kb, demo_corpus, use_llm=False)
    ans = a.answer("What parameters affect porosity in LPBF Ti-6Al-4V?")
    assert ans.mode == "graph_only"
    assert {e.name for e in ans.linked_entities} >= {"porosity", "LPBF", "Ti-6Al-4V"}
    assert "laser power INFLUENCES porosity" in ans.answer
    assert any(it.claim and it.claim["subject"] == "laser power" for it in ans.evidence)
    assert any("SYNTHETIC" in w for w in ans.warnings)


def test_assistant_llm_citations_are_checked(demo_corpus):
    kb = KnowledgeBase.load(demo_corpus)
    llm = FakeLLM(["Laser power influences porosity [E1] [E99].\nInterpretation: higher power may widen the melt pool."])
    ans = ResearchAssistant(kb, demo_corpus, llm=llm).answer("What parameters affect porosity in LPBF Ti-6Al-4V?")
    assert ans.mode == "llm" and ans.cited_ids == ["E1"] and "[E99]" not in ans.answer
    assert ans.interpretation.startswith("higher power")                   # kept separate, labelled
    assert "E99" in " ".join(ans.warnings)
    ans2 = ResearchAssistant(kb, demo_corpus, llm=FakeLLM(["No citations here."])).answer("laser power porosity")
    assert ans2.mode == "graph_only"                                       # uncited answer rejected


def test_gaps_wording(demo_corpus):
    kb = KnowledgeBase.load(demo_corpus)
    rep = find_gaps(kb)
    assert rep.candidates and rep.corpus["papers"] == 5
    for c in rep.candidates:
        assert c.statement.startswith("Limited evidence was found in the current corpus")
        assert "research gap" not in c.statement.lower()
    # parameters are only proposed for processes that report them (no 'laser power' for FDM)
    assert not [c for c in rep.candidates if c.process == "FDM" and c.factor == "laser power"]


def test_evaluation(demo_corpus, lexicon):
    kb = KnowledgeBase.load(demo_corpus)
    rep = evaluate(kb, demo_corpus.path("annotations_dir"), demo_corpus.path("processed_dir"), EntityNormalizer(lexicon))
    assert rep["entities"]["f1"] > 0.8 and rep["relations_strict"]["f1"] > 0.8
    assert rep["normalization"]["total"] > 10 and rep["evidence"]["accuracy"] > 0.8
    q = graph_quality(kb, demo_corpus.path("processed_dir"))
    assert q["unsupported_relation_rate"] == 0.0


def test_chroma_backend_optional(demo_corpus):
    import pytest
    pytest.importorskip("chromadb")
    import hashlib

    def fake_embed(texts):  # deterministic bag-of-words hashing embedding (no model needed)
        out = []
        for t in texts:
            v = [0.0] * 64
            for w in t.lower().split():
                v[int(hashlib.md5(w.encode()).hexdigest(), 16) % 64] += 1.0
            out.append(v)
        return out
    demo_corpus.retrieval.vector_backend = "chroma"
    kb = KnowledgeBase.load(demo_corpus)
    a = ResearchAssistant(kb, demo_corpus, use_llm=False, embed_fn=fake_embed)
    assert a.index.__class__.__name__ == "ChromaIndex"
    hits = a.index.search("laser power porosity", k=3)
    assert hits and all(p.paper_id in kb.papers for p, _ in hits)


def test_export_tables(demo_corpus, tmp_path):
    from openpyxl import load_workbook

    from src.reporting.tables import build_tables, write_csvs, write_workbook
    kb = KnowledgeBase.load(demo_corpus)
    t = build_tables(kb)
    assert len(t["Papers"]) == 5 and t["Property Values"] and t["Parameter Values"] and t["Characterization Methods"]
    row = next(r for r in t["Property Values"] if r["Property"] == "tensile strength" and r["Value"] == 950)
    assert row["Normalized unit"] == "MPa" and row["Evidence (verbatim from paper)"] and row["Page"]
    wb = load_workbook(write_workbook(t, tmp_path / "x.xlsx"))
    assert wb.sheetnames[0] == "Summary" and "Property Values" in wb.sheetnames
    assert str(wb["Summary"]["E4"].value).startswith("=COUNTIF(")
    assert len(write_csvs(t, tmp_path / "csv")) == len(t)
