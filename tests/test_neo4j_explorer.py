"""Neo4j Explorer helpers (no server needed: fake driver and fake graph objects)."""
from src.knowledge_graph.neo4j_explorer import DEMO_QUERIES, PICKER_QUERIES, run_read, to_graph, to_table


class Node:
    def __init__(self, eid, labels, **props):
        self.element_id, self.labels, self._p = eid, set(labels), props

    def items(self):
        return self._p.items()


class Rel:
    def __init__(self, eid, typ, a, b, **props):
        self.element_id, self.type, self.start_node, self.end_node, self._p = eid, typ, a, b, props

    def items(self):
        return self._p.items()


class Path:
    def __init__(self, nodes, rels):
        self.nodes, self.relationships = nodes, rels


LP = Node("n1", ["Entity", "Parameter"], canonical_name="laser power", aliases=["P", "power"])
PO = Node("n2", ["Entity", "Defect"], canonical_name="porosity")
PAPER = Node("n3", ["Paper"], source_file="ti64_lee2025.pdf", title="LPBF Ti-6Al-4V")
CLAIM = Node("n4", ["Claim"], relation="INFLUENCES", evidence="Higher laser power reduced porosity.", page=4)
R = Rel("r1", "INFLUENCES", LP, PO, paper_ids=["p1", "p2"], status="accepted", max_confidence=0.9)


def test_to_graph_collects_nodes_edges_paths_and_lists():
    path = Path([PAPER, CLAIM, LP], [Rel("r2", "SUPPORTS", PAPER, CLAIM), Rel("r3", "SUBJECT", CLAIM, LP)])
    g = to_graph([[LP, R, PO], [path], [[PO, 3]]])
    ids = {n["id"] for n in g["nodes"]}
    assert ids == {"n1", "n2", "n3", "n4"}
    labels = {n["id"]: n["label"] for n in g["nodes"]}
    assert labels["n1"] == "laser power" and labels["n3"] == "ti64_lee2025.pdf" and labels["n4"] == "INFLUENCES"
    infl = next(e for e in g["edges"] if e["id"] == "r1")
    assert infl["papers"] == 2 and "papers: 2" in infl["title"] and infl["from"] == "n1"
    claim = next(n for n in g["nodes"] if n["id"] == "n4")
    assert "Higher laser power reduced porosity." in claim["title"]


def test_to_table_turns_graph_objects_into_text():
    rows = to_table(["a", "r", "b", "n"], [[LP, R, PO, 5]])
    assert rows == [{"a": "laser power", "r": "INFLUENCES", "b": "porosity", "n": 5}]


class FakeResult:
    def __init__(self, keys, rows):
        self._k, self._rows = keys, rows

    def keys(self):
        return self._k

    def __iter__(self):
        return iter([type("Rec", (), {"values": (lambda self, r=r: r)})() for r in self._rows])


class FakeTx:
    def __init__(self, log):
        self.log = log

    def run(self, cypher, **params):
        self.log.append((cypher, params))
        return FakeResult(["x"], [[1], [2], [3]])


class FakeSession:
    def __init__(self, log):
        self.log = log

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def execute_read(self, work):
        self.log.append("READ")
        return work(FakeTx(self.log))


class FakeDriver:
    def __init__(self):
        self.log = []

    def session(self, database=None):
        return FakeSession(self.log)


def test_run_read_uses_a_read_transaction_and_default_params():
    d = FakeDriver()
    keys, rows = run_read(d, "MATCH (n) RETURN n", {"x": 1}, max_rows=2)
    assert d.log[0] == "READ"
    cypher, params = d.log[1]
    assert params["statuses"] == ["accepted"] and "INFLUENCES" in params["effect_rels"] and params["x"] == 1
    assert keys == ["x"] and rows == [[1], [2]]


def test_demo_queries_are_read_only_and_pickers_exist():
    for q in DEMO_QUERIES:
        up = q.cypher.upper()
        assert not any(w in up for w in ("CREATE ", "MERGE ", "DELETE ", "SET ", "REMOVE ")), q.name
        assert q.picker is None or q.picker in PICKER_QUERIES
        assert q.view in ("graph", "table")


def test_excel_export_survives_pdf_control_characters(tmp_path):
    from openpyxl import load_workbook

    from src.reporting.tables import _excel_safe, write_workbook
    assert _excel_safe("300 \x04C") == "300 C" and _excel_safe("= 5") == " = 5" and _excel_safe(3) == 3
    tables = {"Papers": [{"Paper ID": "p1", "Title": "T\x03itle", "Year": 2023, "DEMO / SYNTHETIC": "no"}]}
    for name in ("Property Values", "Parameter Values", "Process Parameters", "Characterization Methods",
                 "Relations & Effects", "Findings", "Entities"):
        tables[name] = [{"Paper ID": "p1", "Evidence (verbatim from paper)": "Table 3 300 \x04C \x03 7.8 mm"}]
    out = write_workbook(tables, tmp_path / "x.xlsx")
    assert load_workbook(out)
