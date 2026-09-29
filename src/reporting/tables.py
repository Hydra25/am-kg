"""Export the extracted knowledge as human-readable tables (Excel + CSV).

One workbook, one sheet per kind of information, every row traceable to
paper -> page -> section -> evidence sentence. Intended for supervisors,
reviewers and thesis appendices.

    python main.py export                      # accepted + review claims
    python main.py export --include-rejected   # also show what QC rejected
    python main.py export --paper paper_6645294073
"""
from __future__ import annotations

import csv
from collections import Counter, defaultdict
from pathlib import Path
from typing import Iterable, Optional

from ..knowledge_graph.kb import KnowledgeBase
from ..ontology.models import Relation

PARAM_TYPES = {"Parameter", "ExperimentalCondition"}
METHOD_TYPE = "CharacterizationMethod"


def _paper_ctx(kb: KnowledgeBase, pid: str) -> tuple[str, str]:
    ctx = kb.paper_context(pid, ("accepted", "review"))
    mats = ", ".join(sorted(kb.entity_name(m) for m in ctx["materials"]))
    procs = ", ".join(sorted(kb.entity_name(p) for p in ctx["processes"]))
    return mats, procs


def _ev(c: Relation) -> dict:
    ev = c.evidence
    return {"Page": ev.page if ev else None, "Section": ev.section if ev else None,
            "Source": ev.source_type if ev else None,
            "Evidence (verbatim from paper)": ev.text if ev else None}


def _qc(c: Relation) -> dict:
    return {"Status": c.status.value, "Confidence": round(c.confidence, 2),
            "Knowledge type": c.knowledge_type.value, "Extracted by": c.method,
            "QC notes": "; ".join(c.qc_reasons), "Claim ID": c.claim_id}


def build_tables(kb: KnowledgeBase, statuses: Iterable[str] = ("accepted", "review"),
                 paper_ids: Optional[set[str]] = None, include_synthetic: bool = True) -> dict[str, list[dict]]:
    st = set(statuses)
    papers = {pid: p for pid, p in kb.papers.items()
              if (paper_ids is None or pid in paper_ids) and (include_synthetic or not p.is_synthetic)}
    claims = [c for c in kb.claims if c.evidence and c.evidence.paper_id in papers and c.status.value in st]
    short = {pid: (p.title or pid)[:70] for pid, p in papers.items()}
    ctx = {pid: _paper_ctx(kb, pid) for pid in papers}

    # ---- Papers ------------------------------------------------------------
    t_papers = [{"Paper ID": pid, "Title": p.title, "Authors": "; ".join(p.authors), "Year": p.year,
                 "Journal": p.journal, "DOI": p.doi, "Source file": p.source_file,
                 "Materials studied": ctx[pid][0], "Processes used": ctx[pid][1],
                 "DEMO / SYNTHETIC": "YES" if p.is_synthetic else "no"}
                for pid, p in sorted(papers.items())]

    # ---- Values (properties / parameters) ------------------------------------
    def value_row(c: Relation, label: str) -> dict:
        q = c.object.quantity
        vctx = c.attributes.get("value_context") or {}
        pid = c.evidence.paper_id  # type: ignore[union-attr]
        return {"Paper ID": pid, "Paper": short[pid], "Materials in paper": ctx[pid][0],
                "Processes in paper": ctx[pid][1], "Processes in sentence": ", ".join(vctx.get("processes") or []),
                label: c.subject.name, "Type": c.subject.type,
                "Value as written": c.object.name,
                "Value": q.value if q else None, "Min": q.value_min if q else None, "Max": q.value_max if q else None,
                "±": q.uncertainty if q else None, "Unit": q.unit if q else None,
                "Normalized value": q.normalized_value if q else None,
                "Normalized min": q.normalized_min if q else None, "Normalized max": q.normalized_max if q else None,
                "Normalized unit": q.normalized_unit if q else None,
                "Material state": vctx.get("state"), **_ev(c), **_qc(c)}

    values = [c for c in claims if c.relation == "HAS_VALUE"]
    t_props = [value_row(c, "Property") for c in values if c.subject.type not in PARAM_TYPES]
    t_params = [value_row(c, "Parameter") for c in values if c.subject.type in PARAM_TYPES]

    # ---- Parameters reported per process (without value) ----------------------
    t_proc_params = []
    agg: dict[tuple, list[Relation]] = defaultdict(list)
    for c in claims:
        if c.relation in ("HAS_PARAMETER", "HAS_CONDITION"):
            agg[(c.evidence.paper_id, c.subject.name, c.object.name, c.object.type)].append(c)  # type: ignore[union-attr]
    for (pid, proc, par, typ), cs in sorted(agg.items()):
        best = max(cs, key=lambda x: x.confidence)
        t_proc_params.append({"Paper ID": pid, "Paper": short[pid], "Process": proc, "Parameter / condition": par,
                              "Type": typ, "Times reported": len(cs), **_ev(best), **_qc(best)})

    # ---- Characterization methods ---------------------------------------------
    t_methods = []
    targets: dict[tuple, set[str]] = defaultdict(set)
    tclaim: dict[tuple, Relation] = {}
    for c in claims:
        if c.relation == "CHARACTERIZES":
            k = (c.evidence.paper_id, c.subject.name)  # type: ignore[union-attr]
            targets[k].add(c.object.name)
            tclaim.setdefault(k, c)
    mentions: Counter = Counter()
    for pid in papers:
        for cid, n in kb.mention_counts.get(pid, {}).items():
            e = kb.entities.get(cid)
            if e and e.entity_type == METHOD_TYPE:
                mentions[(pid, e.canonical_name)] = n
    for (pid, name) in sorted(set(mentions) | set(targets)):
        c = tclaim.get((pid, name))
        t_methods.append({"Paper ID": pid, "Paper": short[pid], "Method": name,
                          "Mentions in paper": mentions.get((pid, name), 0),
                          "Used to characterize": ", ".join(sorted(targets.get((pid, name), []))),
                          **(_ev(c) if c else {"Page": None, "Section": None, "Source": None,
                                               "Evidence (verbatim from paper)": None}),
                          "Status": c.status.value if c else "mention only"})

    # ---- Effect relations (the scientific findings as triples) ------------------
    t_effects = []
    for c in claims:
        if c.relation in ("AFFECTS", "INFLUENCES", "CAUSES", "IMPROVES", "CHANGES", "PRODUCES", "PROCESSES", "FABRICATES"):
            pid = c.evidence.paper_id  # type: ignore[union-attr]
            a = c.attributes
            t_effects.append({"Paper ID": pid, "Paper": short[pid], "Subject": c.subject.name,
                              "Subject type": c.subject.type, "Relation": c.relation, "Object": c.object.name,
                              "Object type": c.object.type, "Trend": a.get("trend"),
                              "No effect (negated)": "YES" if a.get("negated") else "",
                              "Hedged (may/might)": "YES" if a.get("hedged") else "",
                              "From introduction / cited work": "YES" if a.get("claim_origin") == "background" else "",
                              **_ev(c), **_qc(c)})

    # ---- Entities -----------------------------------------------------------------
    t_entities = []
    for pid in sorted(papers):
        for cid, n in kb.mention_counts.get(pid, {}).most_common():
            e = kb.entities.get(cid)
            if not e:
                continue
            t_entities.append({"Paper ID": pid, "Paper": short[pid], "Entity type": e.entity_type,
                               "Canonical name": e.canonical_name, "Mentions": n,
                               "Spellings seen (aliases)": "; ".join(e.aliases[:8]),
                               "In lexicon": "yes" if e.in_lexicon else "NEW - check",
                               "Normalization confidence": round(e.normalization_confidence, 2)})

    # ---- Findings -------------------------------------------------------------------
    t_findings = [{"Paper ID": f.evidence.paper_id, "Paper": short[f.evidence.paper_id], "Finding type": f.finding_type,
                   "Finding (verbatim sentence)": f.text, "Page": f.evidence.page, "Section": f.evidence.section,
                   "Status": f.status.value, "Confidence": round(f.confidence, 2), "Extracted by": f.method}
                  for f in kb.findings if f.evidence.paper_id in papers and f.status.value in st]

    return {"Papers": t_papers, "Property Values": t_props, "Parameter Values": t_params,
            "Process Parameters": t_proc_params, "Characterization Methods": t_methods,
            "Relations & Effects": t_effects, "Entities": t_entities, "Findings": t_findings}


LEGEND = [
    ("How to read this workbook", ""),
    ("Every row", "comes from one paper and shows the page, section and the verbatim evidence sentence (or table row)."),
    ("Status = accepted", "passed all automatic quality checks (confidence >= accept threshold)."),
    ("Status = review", "plausible but needs a human check (new term, cited work, low confidence, derived)."),
    ("Status = rejected", "failed a check (not in source text, invalid relation type, ...). Shown only with --include-rejected."),
    ("Knowledge type EXPLICIT", "the sentence itself states the relation (e.g. 'laser power reduced porosity')."),
    ("Knowledge type EXTRACTED", "built from terms that appear together in one sentence/table row (e.g. a value)."),
    ("Knowledge type DERIVED", "computed by the system (paper-level summary, trend over table rows)."),
    ("Normalized value/unit", "converted to a standard unit (W, mm/s, μm, °C, MPa, J/mm³, %) - original value kept."),
    ("Extracted by", "rule = dictionary + grammar rules; llm = local Qwen model (verified); rule+llm = both agreed."),
    ("DEMO / SYNTHETIC", "papers marked YES are test documents, NOT real research."),
]


def write_workbook(tables: dict[str, list[dict]], out_path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E78")
    fills = {"accepted": PatternFill("solid", fgColor="E2EFDA"), "review": PatternFill("solid", fgColor="FFF2CC"),
             "rejected": PatternFill("solid", fgColor="F8CBAD")}
    wide = {"Evidence (verbatim from paper)": 70, "Finding (verbatim sentence)": 80, "Title": 50, "Paper": 35,
            "QC notes": 30, "Spellings seen (aliases)": 40}

    # --- Summary sheet with live formulas over the other sheets -------------------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Extracted knowledge per paper"
    ws["A1"].font = Font(name="Arial", size=12, bold=True)
    heads = ["Paper ID", "Title", "Year", "DEMO / SYNTHETIC", "Property values", "Parameter values",
             "Process parameters", "Characterization methods", "Relations & effects", "Findings", "Entities"]
    sheets_for = {"Property values": "Property Values", "Parameter values": "Parameter Values",
                  "Process parameters": "Process Parameters", "Characterization methods": "Characterization Methods",
                  "Relations & effects": "Relations & Effects", "Findings": "Findings", "Entities": "Entities"}
    for j, h in enumerate(heads, start=1):
        c = ws.cell(row=3, column=j, value=h)
        c.font, c.fill = bold, head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    n = len(tables["Papers"])
    for i, p in enumerate(tables["Papers"], start=4):
        ws.cell(row=i, column=1, value=p["Paper ID"])
        ws.cell(row=i, column=2, value=p["Title"])
        ws.cell(row=i, column=3, value=p["Year"])
        ws.cell(row=i, column=4, value=p["DEMO / SYNTHETIC"])
        for j, h in enumerate(heads[4:], start=5):
            sheet = sheets_for[h]
            ws.cell(row=i, column=j, value=f"=COUNTIF('{sheet}'!$A:$A,$A{i})")
    total_row = 4 + n
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(name="Arial", size=10, bold=True)
    for j in range(5, len(heads) + 1):
        col = get_column_letter(j)
        ws.cell(row=total_row, column=j, value=f"=SUM({col}4:{col}{total_row - 1})" if n else 0).font = \
            Font(name="Arial", size=10, bold=True)
    for row in ws.iter_rows(min_row=4, max_row=total_row):
        for c in row:
            if not c.font.bold:
                c.font = font
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 60
    for j in range(3, len(heads) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14
    ws.freeze_panes = "A4"
    r = total_row + 2
    for k, v in LEGEND:
        ws.cell(row=r, column=1, value=k).font = Font(name="Arial", size=10, bold=True)
        ws.cell(row=r, column=2, value=v).font = font
        r += 1

    # --- data sheets ----------------------------------------------------------------
    for name, rows in tables.items():
        ws = wb.create_sheet(name)
        cols = list(rows[0].keys()) if rows else ["Paper ID", "(no rows)"]
        for j, h in enumerate(cols, start=1):
            c = ws.cell(row=1, column=j, value=h)
            c.font, c.fill = bold, head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(j)].width = wide.get(h, max(12, min(24, len(h) + 4)))
        status_col = cols.index("Status") + 1 if "Status" in cols else None
        for i, row in enumerate(rows, start=2):
            for j, h in enumerate(cols, start=1):
                v = row.get(h)
                if isinstance(v, (list, dict)):
                    v = str(v)
                c = ws.cell(row=i, column=j, value=v)
                c.font = font
                if h in wide:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            if status_col:
                f = fills.get(str(row.get("Status")))
                if f:
                    ws.cell(row=i, column=status_col).fill = f
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def write_csvs(tables: dict[str, list[dict]], out_dir: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for name, rows in tables.items():
        p = out_dir / (name.lower().replace(" & ", "_").replace(" ", "_") + ".csv")
        cols = list(rows[0].keys()) if rows else []
        with p.open("w", newline="", encoding="utf-8-sig") as fh:  # utf-8-sig -> Excel shows μ, °, ± correctly
            w = csv.DictWriter(fh, fieldnames=cols)
            w.writeheader()
            w.writerows(rows)
        paths.append(p)
    return paths
